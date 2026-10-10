"""Download public original sources; never bypass access controls."""
import argparse
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.deps'))
import requests
from bs4 import BeautifulSoup
import pymupdf
from openpyxl import load_workbook

def collect(refresh=False):
    manifest = json.loads((ROOT/'scripts/source_manifest.json').read_text(encoding='utf-8'))
    previous = {s['source_id']: s for s in json.loads((ROOT/'data/raw/retrieval_log.json').read_text())} if (ROOT/'data/raw/retrieval_log.json').exists() else {}
    logs = []
    for source in manifest:
        row = dict(source)
        target = ROOT/'data/raw'/source['filename']
        row['access_date'] = datetime.now(timezone.utc).date().isoformat()
        try:
            if target.exists() and not refresh:
                body = target.read_bytes()
                row['access_date'] = previous.get(source['source_id'], {}).get('access_date', row['access_date'])
                row['final_url'] = previous.get(source['source_id'], {}).get('final_url', source['source_url'])
                row['http_status'] = previous.get(source['source_id'], {}).get('http_status', '')
            else:
                response = requests.get(source['source_url'], timeout=(15, 60), headers={'User-Agent': 'WasteWiseResearch/1.0 (public scientific data collection)'})
                row.update(http_status=response.status_code, final_url=response.url)
                response.raise_for_status()
                body = response.content
                if source['filename'].endswith('.pdf') and not body.startswith(b'%PDF'):
                    raise ValueError('Expected PDF; received non-PDF content (possible access/error page)')
                if source['filename'].endswith('.html'):
                    lower = body.lower()
                    if b'captcha' in lower and (b'challenge' in lower or b'access denied' in lower):
                        raise ValueError('Access-control challenge; collection stopped for this source')
                target.write_bytes(body)
            if source['filename'].endswith('.xlsx') and not body.startswith(b'PK'):
                failure=target.with_suffix('.failed_response.html')
                if target.exists():
                    target.replace(failure)
                else:
                    failure.write_bytes(body)
                row['failure_response_file']=failure.relative_to(ROOT).as_posix()
                row['retrieval_status']='failed'
                raise ValueError('XLSX download returned a login/access page instead of a spreadsheet; no login or access-control bypass attempted')
            row['sha256'] = hashlib.sha256(body).hexdigest()
            row['bytes'] = len(body)
            if source['filename'].endswith('.pdf'):
                with pymupdf.open(target) as document:
                    for i, page in enumerate(document, 1):
                        extracted = ROOT/'data/processed'/f"{source['source_id']}_page_{i:03}.txt"
                        # Use a single extractor consistently, including cached downloads.
                        extracted.write_text(page.get_text(sort=True) or '', encoding='utf-8')
                    row['page_count'] = len(document)
            elif source['filename'].endswith('.xlsx'):
                workbook=load_workbook(target,read_only=True,data_only=True)
                lines=[]
                for sheet in workbook:
                    lines.append('SHEET: '+sheet.title)
                    for values in sheet.iter_rows(values_only=True):
                        lines.append('\t'.join('' if v is None else str(v) for v in values))
                workbook.close()
                (ROOT/'data/processed'/f"{source['source_id']}.txt").write_text('\n'.join(lines),encoding='utf-8')
            else:
                soup = BeautifulSoup(body, 'html.parser')
                for element in soup(['script','style','nav','footer','header']):
                    element.decompose()
                text = (soup.find('main') or soup).get_text(' ', strip=True)
                (ROOT/'data/processed'/f"{source['source_id']}.txt").write_text(text, encoding='utf-8')
            row['retrieval_status'] = 'downloaded'
        except Exception as exc:
            row['retrieval_status'] = 'failed'
            row['error'] = f'{type(exc).__name__}: {exc}'
        logs.append(row)
        print(source['source_id'], row['retrieval_status'], row.get('bytes', ''), row.get('error', '')[:180], flush=True)
        (ROOT/'data/raw/retrieval_log.json').write_text(json.dumps(logs, indent=2), encoding='utf-8')
    return logs

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--refresh', action='store_true', help='Redownload sources; default reuses archived originals')
    args = parser.parse_args()
    collect(args.refresh)
